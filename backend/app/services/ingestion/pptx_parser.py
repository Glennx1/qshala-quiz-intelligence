import os
import re
import uuid
import zipfile
import logging
import mimetypes
from pathlib import Path
from typing import List, Dict, Any, Optional, Set
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from backend.app.config import settings

logger = logging.getLogger(__name__)

AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".aac", ".wma", ".ogg", ".flac"}
VIDEO_EXTS = {".mp4", ".mov", ".avi", ".wmv", ".mkv", ".webm", ".m4v"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tiff"}

class PPTXParser:
    """
    Extracts structured content, text, speaker notes, tables,
    and embedded multi-modal media (images, audio clips, video clips)
    from PPTX presentation decks using python-pptx shapes, slide.part.rels,
    and zipfile package fallback.
    """
    def __init__(self, file_path: str, document_id: str, write_to_disk: bool = True):
        self.file_path = file_path
        self.document_id = document_id
        self.write_to_disk = write_to_disk
        self.output_dir = settings.SLIDES_DIR / document_id
        self.output_media_dir = self.output_dir / "media"
        if self.write_to_disk:
            self.output_dir.mkdir(parents=True, exist_ok=True)
            self.output_media_dir.mkdir(parents=True, exist_ok=True)

    def _classify_media_type(self, ext: str) -> str:
        ext_lower = ext.lower()
        if ext_lower in AUDIO_EXTS:
            return "audio"
        if ext_lower in VIDEO_EXTS:
            return "video"
        if ext_lower in IMAGE_EXTS:
            return "image"
        return "other"

    def extract_media_from_zip(self) -> Dict[int, List[Dict[str, Any]]]:
        """
        Fallback & supplementary extractor: inspects the PPTX zip archive directly,
        parsing ppt/slides/_rels/slide{N}.xml.rels to map ppt/media/* files to slides.
        """
        slide_media_map: Dict[int, List[Dict[str, Any]]] = {}
        if not os.path.exists(self.file_path):
            return slide_media_map

        try:
            with zipfile.ZipFile(self.file_path, "r") as z:
                # Find all media files in archive
                media_entries = {
                    name: z.read(name)
                    for name in z.namelist()
                    if name.startswith("ppt/media/") and not name.endswith("/")
                }

                if not media_entries:
                    return slide_media_map

                # Map relationships: ppt/slides/_rels/slide{N}.xml.rels -> target
                slide_rel_pattern = re.compile(r"ppt/slides/_rels/slide(\d+)\.xml\.rels")
                mapped_media_keys = set()

                for rel_name in z.namelist():
                    m = slide_rel_pattern.match(rel_name)
                    if not m:
                        continue
                    slide_num = int(m.group(1))
                    rel_xml = z.read(rel_name).decode("utf-8", errors="ignore")

                    # Extract target references
                    targets = re.findall(r'Target="([^"]+)"', rel_xml)
                    for target in targets:
                        clean_target = target.replace("../", "ppt/")
                        if clean_target in media_entries:
                            mapped_media_keys.add(clean_target)
                            data = media_entries[clean_target]
                            ext = Path(clean_target).suffix.lower()
                            media_type = self._classify_media_type(ext)
                            mime, _ = mimetypes.guess_type(clean_target)

                            item = {
                                "filename": Path(clean_target).name,
                                "slide_number": slide_num,
                                "media_type": media_type,
                                "raw_bytes": data,
                                "mime_type": mime or "application/octet-stream",
                                "source": "zip_rel"
                            }
                            slide_media_map.setdefault(slide_num, []).append(item)

                # Any unmapped media is placed under slide 1 as general presentation media
                for key, data in media_entries.items():
                    if key not in mapped_media_keys:
                        ext = Path(key).suffix.lower()
                        media_type = self._classify_media_type(ext)
                        mime, _ = mimetypes.guess_type(key)
                        item = {
                            "filename": Path(key).name,
                            "slide_number": 1,
                            "media_type": media_type,
                            "raw_bytes": data,
                            "mime_type": mime or "application/octet-stream",
                            "source": "zip_unmapped"
                        }
                        slide_media_map.setdefault(1, []).append(item)

        except Exception as e:
            logger.warning(f"PPTX zipfile fallback inspection warning: {e}")

        return slide_media_map

    def parse(self) -> List[Dict[str, Any]]:
        prs = Presentation(self.file_path)
        slides_data = []
        zip_media = self.extract_media_from_zip()

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_text_chunks = []
            title = None
            speaker_notes = ""
            image_paths = []
            audio_items = []
            video_items = []
            all_media_items = []
            seen_media_names: Set[str] = set()

            # Extract title if present
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()

            # Extract speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                speaker_notes = slide.notes_slide.notes_text_frame.text.strip()

            # 1. Iterate shapes for text, tables, and PICTURE shapes
            img_counter = 1
            for shape in slide.shapes:
                if shape.has_text_frame:
                    shape_text = shape.text_frame.text.strip()
                    if shape_text:
                        slide_text_chunks.append(shape_text)

                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        slide_text_chunks.append("\n".join(table_rows))

                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        ext = f".{image.ext.lower().lstrip('.')}"
                        img_filename = f"slide_{slide_idx}_img_{img_counter}{ext}"
                        img_path = self.output_media_dir / img_filename
                        if self.write_to_disk:
                            with open(img_path, "wb") as f:
                                f.write(image.blob)

                        rel_path = f"/static/slides/{self.document_id}/media/{img_filename}"
                        image_paths.append(rel_path)
                        seen_media_names.add(img_filename)

                        all_media_items.append({
                            "filename": img_filename,
                            "slide_number": slide_idx,
                            "media_type": "image",
                            "raw_bytes": image.blob,
                            "mime_type": f"image/{image.ext}",
                            "local_path": str(img_path) if self.write_to_disk else None,
                            "rel_url": rel_path
                        })
                        img_counter += 1
                    except Exception as e:
                        logger.warning(f"Could not extract picture shape from slide {slide_idx}: {e}")

            # 2. Extract embedded audio & video from slide.part.rels
            try:
                for rel_id, rel in slide.part.rels.items():
                    rel_type = (rel.reltype or "").lower()
                    target_ref = getattr(rel, "target_ref", "") or ""
                    ext = Path(target_ref).suffix.lower()

                    is_media = (
                        "video" in rel_type or
                        "audio" in rel_type or
                        "media" in rel_type or
                        ext in AUDIO_EXTS or
                        ext in VIDEO_EXTS or
                        ext in IMAGE_EXTS
                    )

                    if is_media and hasattr(rel, "target_part"):
                        target_part = rel.target_part
                        raw_data = getattr(target_part, "blob", None)
                        if raw_data:
                            part_name = Path(target_part.partname).name if hasattr(target_part, "partname") else f"rel_{rel_id}{ext}"
                            if part_name not in seen_media_names:
                                seen_media_names.add(part_name)
                                media_type = self._classify_media_type(ext)
                                mime_type = getattr(target_part, "content_type", None) or mimetypes.guess_type(part_name)[0] or "application/octet-stream"

                                local_file = self.output_media_dir / part_name
                                if self.write_to_disk:
                                    with open(local_file, "wb") as f:
                                        f.write(raw_data)

                                rel_url = f"/static/slides/{self.document_id}/media/{part_name}"
                                media_dict = {
                                    "filename": part_name,
                                    "slide_number": slide_idx,
                                    "media_type": media_type,
                                    "raw_bytes": raw_data,
                                    "mime_type": mime_type,
                                    "local_path": str(local_file) if self.write_to_disk else None,
                                    "rel_url": rel_url
                                }
                                all_media_items.append(media_dict)
                                if media_type == "audio":
                                    audio_items.append(media_dict)
                                elif media_type == "video":
                                    video_items.append(media_dict)
                                elif media_type == "image":
                                    image_paths.append(rel_url)
            except Exception as re_err:
                logger.warning(f"Error inspecting slide.part.rels on slide {slide_idx}: {re_err}")

            # 3. Incorporate any media found from the zipfile inspection for this slide
            if slide_idx in zip_media:
                for z_item in zip_media[slide_idx]:
                    if z_item["filename"] not in seen_media_names:
                        seen_media_names.add(z_item["filename"])
                        local_file = self.output_media_dir / z_item["filename"]
                        if self.write_to_disk:
                            with open(local_file, "wb") as f:
                                f.write(z_item["raw_bytes"])

                        rel_url = f"/static/slides/{self.document_id}/media/{z_item['filename']}"
                        item_record = {
                            "filename": z_item["filename"],
                            "slide_number": slide_idx,
                            "media_type": z_item["media_type"],
                            "raw_bytes": z_item["raw_bytes"],
                            "mime_type": z_item["mime_type"],
                            "local_path": str(local_file) if self.write_to_disk else None,
                            "rel_url": rel_url
                        }
                        all_media_items.append(item_record)
                        if z_item["media_type"] == "audio":
                            audio_items.append(item_record)
                        elif z_item["media_type"] == "video":
                            video_items.append(item_record)
                        elif z_item["media_type"] == "image":
                            image_paths.append(rel_url)

            # Deduplicate text chunks while preserving order
            unique_chunks = []
            seen = set()
            for chunk in slide_text_chunks:
                norm = chunk.strip()
                if norm and norm not in seen:
                    unique_chunks.append(norm)
                    seen.add(norm)

            combined_text = "\n\n".join(unique_chunks)

            slides_data.append({
                "slide_number": slide_idx,
                "title": title or (unique_chunks[0] if unique_chunks else f"Slide {slide_idx}"),
                "extracted_text": combined_text,
                "speaker_notes": speaker_notes,
                "image_paths": image_paths,
                "has_images": len(image_paths) > 0,
                "audio_items": audio_items,
                "video_items": video_items,
                "all_media_items": all_media_items,
                "metadata_json": {
                    "shape_count": len(slide.shapes),
                    "has_notes": bool(speaker_notes),
                    "audio_count": len(audio_items),
                    "video_count": len(video_items),
                    "image_count": len(image_paths)
                }
            })

        return slides_data

    def extract_all_media(self) -> List[Dict[str, Any]]:
        """
        Parses presentation and aggregates all multi-modal media objects across slides.
        """
        slides = self.parse()
        all_media = []
        for s in slides:
            all_media.extend(s.get("all_media_items", []))
        return all_media
