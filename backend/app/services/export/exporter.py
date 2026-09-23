import csv
import io
import json
from typing import Dict, Any, List
from backend.app.models.quiz import Quiz

class QuizExporter:
    @staticmethod
    def to_json(quiz: Quiz) -> str:
        data = {
            "quiz_id": quiz.id,
            "title": quiz.title,
            "topic": quiz.topic,
            "subtopic": quiz.subtopic,
            "grade_min": quiz.grade_min,
            "grade_max": quiz.grade_max,
            "difficulty": quiz.difficulty,
            "difficulty_distribution": quiz.difficulty_distribution,
            "generation_mode": quiz.generation_mode,
            "created_at": quiz.created_at.isoformat() if quiz.created_at else None,
            "questions": [
                {
                    "order": q.order_index,
                    "question": q.question_text,
                    "options": q.options,
                    "answer": q.answer,
                    "explanation": q.explanation,
                    "difficulty": q.difficulty,
                    "grade_range": f"{q.grade_min}-{q.grade_max}" if q.grade_min else "General",
                    "question_type": q.question_type,
                    "validation_status": q.validation_status,
                    "sources": [
                        {
                            "document_title": src.document_title,
                            "slide_number": src.slide_number,
                            "quote": src.source_quote,
                            "rationale": src.rationale
                        }
                        for src in q.retrieval_sources
                    ]
                }
                for q in quiz.questions
                if q.is_approved
            ]
        }
        return json.dumps(data, indent=2)

    @staticmethod
    def to_csv(quiz: Quiz) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Order",
            "Question",
            "Options",
            "Answer",
            "Explanation",
            "Difficulty",
            "Topic",
            "Primary Source"
        ])

        for q in quiz.questions:
            if not q.is_approved:
                continue

            primary_src = ""
            if q.retrieval_sources:
                s = q.retrieval_sources[0]
                primary_src = f"{s.document_title or 'Deck'} (Slide {s.slide_number or 'N/A'})"

            options_str = " | ".join(q.options) if q.options else "N/A"

            writer.writerow([
                q.order_index,
                q.question_text,
                options_str,
                q.answer,
                q.explanation or "",
                q.difficulty,
                q.topic or quiz.topic,
                primary_src
            ])

        return output.getvalue()

    @staticmethod
    def to_pptx(quiz: Quiz) -> bytes:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE

        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(5.625) # 16:9 widescreen
        blank_layout = prs.slide_layouts[6]

        # Colors
        c_slate_900 = RGBColor(15, 23, 42)
        c_slate_800 = RGBColor(30, 41, 59)
        c_slate_700 = RGBColor(51, 65, 85)
        c_slate_500 = RGBColor(100, 116, 139)
        c_slate_400 = RGBColor(148, 163, 184)
        c_slate_200 = RGBColor(226, 232, 240)
        c_slate_100 = RGBColor(241, 245, 249)
        c_white = RGBColor(255, 255, 255)
        c_blue_600 = RGBColor(37, 99, 235)
        c_blue_500 = RGBColor(59, 130, 246)
        c_emerald_600 = RGBColor(5, 150, 105)
        c_emerald_500 = RGBColor(16, 185, 129)
        c_amber_500 = RGBColor(245, 158, 11)
        c_rose_500 = RGBColor(244, 63, 94)

        # ---------------------------------------------------------------------
        # SLIDE 1: Title Slide (Dark Theme)
        # ---------------------------------------------------------------------
        slide_title = prs.slides.add_slide(blank_layout)

        # Background fill
        bg = slide_title.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_slate_900
        bg.line.fill.background()

        # Top Accent Ribbon
        accent = slide_title.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = c_blue_500
        accent.line.fill.background()

        tb_title = slide_title.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.4), Inches(3.2))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True

        p1 = tf_title.paragraphs[0]
        p1.text = quiz.title
        p1.font.size = Pt(32)
        p1.font.bold = True
        p1.font.color.rgb = c_white

        p2 = tf_title.add_paragraph()
        p2.space_before = Pt(12)
        grade_info = f" · Grades {quiz.grade_min}–{quiz.grade_max}" if quiz.grade_min else ""
        p2.text = f"Topic: {quiz.topic}{grade_info}  |  Difficulty: {quiz.difficulty}"
        p2.font.size = Pt(16)
        p2.font.color.rgb = c_slate_400

        p3 = tf_title.add_paragraph()
        p3.space_before = Pt(24)
        p3.text = "QShala Quiz Intelligence Platform  •  Tournament Master Deck"
        p3.font.size = Pt(12)
        p3.font.color.rgb = c_blue_500
        p3.font.bold = True

        # ---------------------------------------------------------------------
        # Question & Answer Slide Pairs
        # ---------------------------------------------------------------------
        approved_questions = [q for q in quiz.questions if q.is_approved]
        total_q = len(approved_questions)

        for idx, q in enumerate(approved_questions, start=1):
            diff_color = c_emerald_500 if q.difficulty == "Easy" else (c_amber_500 if q.difficulty == "Medium" else c_rose_500)

            # -----------------------------------------------------------------
            # QUESTION SLIDE (Dark Studio Canvas)
            # -----------------------------------------------------------------
            s_q = prs.slides.add_slide(blank_layout)

            # Dark Background
            bg_q = s_q.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
            bg_q.fill.solid()
            bg_q.fill.fore_color.rgb = c_slate_900
            bg_q.line.fill.background()

            # Top bar badge
            tb_q_badge = s_q.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.4), Inches(0.5))
            tf_q_badge = tb_q_badge.text_frame
            p_badge = tf_q_badge.paragraphs[0]
            p_badge.text = f"QUESTION {idx} OF {total_q}   •   {q.difficulty.upper()}"
            p_badge.font.size = Pt(12)
            p_badge.font.bold = True
            p_badge.font.color.rgb = diff_color

            # Question Text
            tb_q_body = s_q.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(2.8))
            tf_q_body = tb_q_body.text_frame
            tf_q_body.word_wrap = True
            p_q_text = tf_q_body.paragraphs[0]
            p_q_text.text = q.question_text
            p_q_text.font.size = Pt(24)
            p_q_text.font.bold = True
            p_q_text.font.color.rgb = c_white

            # Render options if MCQ
            if q.options and len(q.options) >= 2:
                tb_opt = s_q.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8.4), Inches(1.6))
                tf_opt = tb_opt.text_frame
                tf_opt.word_wrap = True
                for opt_idx, opt_text in enumerate(q.options):
                    p_opt = tf_opt.add_paragraph() if opt_idx > 0 else tf_opt.paragraphs[0]
                    p_opt.text = opt_text
                    p_opt.font.size = Pt(15)
                    p_opt.font.color.rgb = c_slate_200
                    p_opt.space_before = Pt(4)

            # Footer
            tb_q_foot = s_q.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(8.4), Inches(0.4))
            p_q_foot = tb_q_foot.text_frame.paragraphs[0]
            p_q_foot.text = f"QShala Tournament  •  {quiz.topic}"
            p_q_foot.font.size = Pt(10)
            p_q_foot.font.color.rgb = c_slate_500

            # -----------------------------------------------------------------
            # ANSWER SLIDE (Crisp Light Canvas with Emerald Accent)
            # -----------------------------------------------------------------
            s_a = prs.slides.add_slide(blank_layout)

            # White Background
            bg_a = s_a.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
            bg_a.fill.solid()
            bg_a.fill.fore_color.rgb = c_white
            bg_a.line.fill.background()

            # Top Accent Bar
            acc_a = s_a.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(0.12))
            acc_a.fill.solid()
            acc_a.fill.fore_color.rgb = c_emerald_600
            acc_a.line.fill.background()

            # Header badge
            tb_a_badge = s_a.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(0.4))
            p_a_badge = tb_a_badge.text_frame.paragraphs[0]
            p_a_badge.text = f"ANSWER {idx} / {total_q}"
            p_a_badge.font.size = Pt(12)
            p_a_badge.font.bold = True
            p_a_badge.font.color.rgb = c_emerald_600

            # Answer Title & Value
            tb_a_main = s_a.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(8.4), Inches(1.2))
            tf_a_main = tb_a_main.text_frame
            tf_a_main.word_wrap = True
            p_ans_val = tf_a_main.paragraphs[0]
            p_ans_val.text = q.answer
            p_ans_val.font.size = Pt(28)
            p_ans_val.font.bold = True
            p_ans_val.font.color.rgb = c_slate_900

            # Explanation Box (Styled Card)
            if q.explanation:
                exp_card = s_a.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.25), Inches(8.4), Inches(2.5))
                exp_card.fill.solid()
                exp_card.fill.fore_color.rgb = c_slate_100
                exp_card.line.color.rgb = c_slate_200
                exp_card.line.width = Pt(1)

                tf_exp = exp_card.text_frame
                tf_exp.word_wrap = True
                tf_exp.margin_left = Inches(0.25)
                tf_exp.margin_top = Inches(0.2)
                tf_exp.margin_right = Inches(0.25)

                p_exp_h = tf_exp.paragraphs[0]
                p_exp_h.text = "THE STORY BEHIND THE ANSWER"
                p_exp_h.font.size = Pt(11)
                p_exp_h.font.bold = True
                p_exp_h.font.color.rgb = c_blue_600

                p_exp_b = tf_exp.add_paragraph()
                p_exp_b.space_before = Pt(6)
                p_exp_b.text = q.explanation
                p_exp_b.font.size = Pt(14)
                p_exp_b.font.color.rgb = c_slate_700

            # Footer / Provenance note
            tb_a_foot = s_a.shapes.add_textbox(Inches(0.8), Inches(4.9), Inches(8.4), Inches(0.4))
            p_a_foot = tb_a_foot.text_frame.paragraphs[0]
            src_doc = "QShala Archive"
            src_slide = "N/A"
            if q.retrieval_sources:
                src = q.retrieval_sources[0]
                src_doc = src.document_title or src_doc
                src_slide = str(src.slide_number) if src.slide_number else src_slide
            p_a_foot.text = f"Source Provenance: {src_doc} (Slide {src_slide})"
            p_a_foot.font.size = Pt(10)
            p_a_foot.font.color.rgb = c_slate_500

            # Speaker notes for presenter
            notes_slide = s_a.notes_slide
            tf_notes = notes_slide.notes_text_frame
            tf_notes.text = f"Question: {q.question_text}\nAnswer: {q.answer}\nSource: {src_doc} (Slide {src_slide})\nNotes: {q.explanation or ''}"

        # ---------------------------------------------------------------------
        # SLIDE FINAL: Summary / Concluding Slide
        # ---------------------------------------------------------------------
        s_end = prs.slides.add_slide(blank_layout)
        bg_end = s_end.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(5.625))
        bg_end.fill.solid()
        bg_end.fill.fore_color.rgb = c_slate_900
        bg_end.line.fill.background()

        tb_end = s_end.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(8.0), Inches(2.5))
        tf_end = tb_end.text_frame
        tf_end.word_wrap = True

        p_end_1 = tf_end.paragraphs[0]
        p_end_1.text = "Quiz Complete!"
        p_end_1.font.size = Pt(36)
        p_end_1.font.bold = True
        p_end_1.font.color.rgb = c_white

        p_end_2 = tf_end.add_paragraph()
        p_end_2.space_before = Pt(10)
        p_end_2.text = f"Compiled {total_q} questions across {quiz.topic}  |  Difficulty: {quiz.difficulty}"
        p_end_2.font.size = Pt(16)
        p_end_2.font.color.rgb = c_slate_400

        p_end_3 = tf_end.add_paragraph()
        p_end_3.space_before = Pt(20)
        p_end_3.text = "Curiosity • Learning • Intelligence — Powered by QShala"
        p_end_3.font.size = Pt(12)
        p_end_3.font.color.rgb = c_emerald_500
        p_end_3.font.bold = True

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
