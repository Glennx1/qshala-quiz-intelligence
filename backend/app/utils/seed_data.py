import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SEED_DIR = Path(__file__).resolve().parent.parent.parent / "sample_decks"
SEED_DIR.mkdir(parents=True, exist_ok=True)

def create_australian_history_deck() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    blank_slide_layout = prs.slide_layouts[6]

    # Slide 1: Title Slide
    slide = prs.slides.add_slide(blank_slide_layout)
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "QShala Quest: Australian History & Explorers"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(20, 40, 80)
    
    p2 = tf.add_paragraph()
    p2.text = "Curated Tournament Round for Grades 3 to 5 (Year: 2022)"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(100, 100, 100)

    # Q&A data (Authentic QShala format: Question Slide -> Next Slide Answer with Explanation)
    qa_pairs = [
        (
            "Who was the British naval officer and navigator who mapped the eastern coastline of Australia aboard HMS Endeavour in 1770?",
            "Captain James Cook",
            "Captain James Cook charted Botany Bay and the east coast in 1770, claiming New South Wales for Britain aboard HMS Endeavour. His voyage provided the detailed charts that led to European settlement 18 years later."
        ),
        (
            "What is the historic name given to the convoy of 11 British ships that arrived in Sydney Cove in January 1788 to establish the first European colony?",
            "The First Fleet",
            "Commanded by Captain Arthur Phillip, the First Fleet carried over 1,400 convicts, marines, and officials across 15,000 miles of ocean over 250 days, arriving on 26 January 1788."
        ),
        (
            "For tens of thousands of years before European arrival, Aboriginal Australians shared oral sacred creation stories known in English by what name?",
            "The Dreamtime (Jukurrpa)",
            "The Dreamtime encompasses the ancient spiritual, cultural, and geographic history of the world created by Ancestral Beings. These sacred stories and songlines map out laws, traditions, and geography across Australia."
        ),
        (
            "On January 1, 1901, the six separate British colonies federated to form the Commonwealth of Australia. Who served as the nation's inaugural Prime Minister?",
            "Sir Edmund Barton",
            "Sir Edmund Barton was sworn in as Australia's inaugural Prime Minister at Centennial Park, Sydney on New Year's Day 1901. He later resigned to become one of the founding justices of the High Court of Australia."
        ),
        (
            "Why was the inland planned city of Canberra chosen as Australia's federal capital in 1913 instead of Sydney or Melbourne?",
            "As a historic compromise between rival cities Sydney and Melbourne",
            "Sydney and Melbourne fiercely contested which should be the national capital. Section 125 of the Australian Constitution established that the capital must be in New South Wales but at least 100 miles from Sydney, leading to Walter Burley Griffin's design of Canberra."
        ),
        (
            "Which legendary Australian bushranger and folk figure wore 44-kilogram homemade iron plate armor during his final standoff in Glenrowan in 1880?",
            "Ned Kelly",
            "Ned Kelly and the Kelly Gang forged heavy protective suits from ploughshares before their famous 1880 siege at Glenrowan. His iconic helmet and suit are now preserved at the State Library of Victoria."
        ),
        (
            "In 1854, gold prospectors in Ballarat, Victoria rose up against oppressive colonial mining licenses in which pivotal rebellion for democratic rights?",
            "The Eureka Stockade",
            "The Eureka Stockade was a pivotal rebellion where miners swore allegiance under the Southern Cross flag against taxation without representation. It is widely regarded as the birthplace of Australian democracy."
        ),
        (
            "Which aerodynamic wooden hunting and ceremonial tool was developed by Indigenous Australians over tens of thousands of years and is world-famous for returning when thrown?",
            "Boomerang",
            "Indigenous Australians mastered sophisticated principles of aerodynamics long before modern flight, designing returning boomerangs for bird decoy hunting, music, and ceremonies."
        ),
        (
            "What iconic sacred sandstone monolith in the Northern Territory desert was officially handed back to its traditional Anangu owners in October 1985?",
            "Uluru (Ayers Rock)",
            "Uluru is deeply sacred to the Anangu people. In a historic handback ceremony on 26 October 1985, the Australian Governor-General presented the title deeds back to the traditional owners."
        ),
        (
            "In 1902, the Commonwealth Franchise Act made Australia one of the earliest modern nations to grant which democratic right to women nationally?",
            "The right to vote and stand for federal parliament",
            "Australia became one of the first countries in the world where women won both the right to vote in federal elections and to stand for federal parliament in 1902, following South Australia's pioneering 1894 legislation."
        )
    ]

    for idx, (q_text, ans_text, expl_text) in enumerate(qa_pairs, start=1):
        # Slide 1: Question Slide
        slide_q = prs.slides.add_slide(blank_slide_layout)
        tb_q = slide_q.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(4.0))
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True
        
        p_qheader = tf_q.paragraphs[0]
        p_qheader.text = f"QUESTION {idx}"
        p_qheader.font.size = Pt(14)
        p_qheader.font.bold = True
        p_qheader.font.color.rgb = RGBColor(100, 116, 139)

        p_q = tf_q.add_paragraph()
        p_q.text = q_text
        p_q.font.size = Pt(24)
        p_q.font.bold = True
        p_q.font.color.rgb = RGBColor(15, 23, 42)

        # Slide 2: Answer & Explanation Slide
        slide_a = prs.slides.add_slide(blank_slide_layout)
        tb_a = slide_a.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(4.0))
        tf_a = tb_a.text_frame
        tf_a.word_wrap = True

        p_aheader = tf_a.paragraphs[0]
        p_aheader.text = f"ANSWER {idx}"
        p_aheader.font.size = Pt(14)
        p_aheader.font.bold = True
        p_aheader.font.color.rgb = RGBColor(16, 185, 129)

        p_ans = tf_a.add_paragraph()
        p_ans.text = f"Answer: {ans_text}"
        p_ans.font.size = Pt(26)
        p_ans.font.bold = True
        p_ans.font.color.rgb = RGBColor(5, 150, 105)

        p_gap = tf_a.add_paragraph()
        p_gap.text = ""

        p_expl_title = tf_a.add_paragraph()
        p_expl_title.text = "Explanation & Context:"
        p_expl_title.font.size = Pt(15)
        p_expl_title.font.bold = True
        p_expl_title.font.color.rgb = RGBColor(71, 85, 105)

        p_expl = tf_a.add_paragraph()
        p_expl.text = expl_text
        p_expl.font.size = Pt(17)
        p_expl.font.color.rgb = RGBColor(51, 65, 85)

    target_path = SEED_DIR / "Australian_History_Quiz_2022.pptx"
    prs.save(target_path)
    return target_path

def create_world_geography_deck() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank_layout = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank_layout)
    tb = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "QShala Junior Explorers: World Wonders & Oceans"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(20, 40, 80)

    qa_list = [
        (
            "Which Australian coral reef system is recognized as the largest living structure on Planet Earth and can be seen from space?",
            "The Great Barrier Reef",
            "Spanning over 2,300 kilometers off the Queensland coast, the Great Barrier Reef is composed of billions of tiny coral polyps and supports thousands of species of marine life."
        ),
        (
            "Which majestic peak in the Himalayas is the highest mountain above sea level on Earth at 8,848.86 meters?",
            "Mount Everest",
            "Mount Everest, called Sagarmatha in Nepal and Chomolungma in Tibet, stands proudly on the border between Nepal and China. It was first officially summitted by Sir Edmund Hillary and Tenzing Norgay in 1953."
        )
    ]

    for idx, (q_text, ans, expl) in enumerate(qa_list, start=1):
        # Slide 1: Question Slide
        slide_q = prs.slides.add_slide(blank_layout)
        tb_q = slide_q.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(4.0))
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True

        p_qheader = tf_q.paragraphs[0]
        p_qheader.text = f"QUESTION {idx}"
        p_qheader.font.size = Pt(14)
        p_qheader.font.bold = True
        p_qheader.font.color.rgb = RGBColor(100, 116, 139)

        p_q = tf_q.add_paragraph()
        p_q.text = q_text
        p_q.font.size = Pt(24)
        p_q.font.bold = True
        p_q.font.color.rgb = RGBColor(15, 23, 42)

        # Slide 2: Answer & Explanation Slide
        slide_a = prs.slides.add_slide(blank_layout)
        tb_a = slide_a.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(4.0))
        tf_a = tb_a.text_frame
        tf_a.word_wrap = True

        p_aheader = tf_a.paragraphs[0]
        p_aheader.text = f"ANSWER {idx}"
        p_aheader.font.size = Pt(14)
        p_aheader.font.bold = True
        p_aheader.font.color.rgb = RGBColor(16, 185, 129)

        pa = tf_a.add_paragraph()
        pa.text = f"Answer: {ans}"
        pa.font.size = Pt(26)
        pa.font.bold = True
        pa.font.color.rgb = RGBColor(5, 150, 105)

        p_gap = tf_a.add_paragraph()
        p_gap.text = ""

        pe_title = tf_a.add_paragraph()
        pe_title.text = "Explanation & Context:"
        pe_title.font.size = Pt(15)
        pe_title.font.bold = True
        pe_title.font.color.rgb = RGBColor(71, 85, 105)

        pe = tf_a.add_paragraph()
        pe.text = expl
        pe.font.size = Pt(17)
        pe.font.color.rgb = RGBColor(51, 65, 85)

    target_path = SEED_DIR / "World_Geography_and_Oceans_2023.pptx"
    prs.save(target_path)
    return target_path

if __name__ == "__main__":
    p1 = create_australian_history_deck()
    p2 = create_world_geography_deck()
    print(f"Generated QShala slide pairs:\n - {p1}\n - {p2}")
