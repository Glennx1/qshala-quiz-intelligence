"""Test that IngestionPipeline recovers from a stale/idle DB connection during dedupe/save."""
import asyncio
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from sqlalchemy.exc import OperationalError

from backend.app.config import settings
from backend.app.database import SessionLocal
from backend.app.models.document import Document
from backend.app.models.question import Question
from backend.app.services.ingestion.pipeline import IngestionPipeline, INGESTION_STATUS_REGISTRY


class FakeGemini:
    """Mock Gemini provider to avoid external API calls during test."""
    async def generate_json_schema(self, prompt, schema, system_instruction=None, model=None,
                                   thinking_level=None, temperature=0.0):
        data = {
            "topics": [
                {"subtopic_l2": "Modern History", "keywords": ["Exploration", "Voyage"]},
            ],
            "geo_location": "Global",
            "india_international": "INTL",
            "relevant_year": 2022,
            "fit": "good",
            "search_summary": "Test question.",
        }
        return data, {"in": 10, "out": 10, "think": 0, "cached": 0}


def _build_test_deck(path: Path, tag: str, n: int = 2):
    prs = Presentation()
    blank = prs.slide_layouts[6]

    def add(text):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
        tb.text_frame.text = text

    add(f"Quiz Header {tag}")
    for i in range(1, n + 1):
        add(f"QUESTION {i}\nWhat is the test question number {i} for {tag}?")
        add(f"ANSWER {i}\nAnswer: Test Answer {tag} {i}\n\nContext for {tag} {i}.")
    prs.save(str(path))


def test_stale_session_recovery_during_save():
    """Simulate a stale connection mid-pipeline (IdleInTransactionSessionTimeout on dedupe)

    Verifies that the retry-once logic catches OperationalError, acquires a fresh SessionLocal,
    and successfully saves all questions and marks the document COMPLETED.
    """
    tag = uuid.uuid4().hex[:8]
    dest = settings.UPLOAD_DIR / f"test_stale_{tag}.pptx"
    settings.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    _build_test_deck(dest, tag, n=2)

    doc_id = str(uuid.uuid4())
    db = SessionLocal()
    try:
        db.add(Document(
            id=doc_id,
            filename=f"2024 Test Stale Quiz {tag}.pptx",
            title=f"Test Stale Quiz {tag}",
            year=2024,
            file_type="pptx",
            storage_path=str(dest),
            file_size_bytes=dest.stat().st_size,
            processing_status="PROCESSING"
        ))
        db.commit()
    finally:
        db.close()

    pipe = IngestionPipeline()
    pipe.qshala_tagger.provider = FakeGemini()
    pipe.qshala_tagger.rpm = 0

    real_deduplicate_batch = pipe.deduplicator.deduplicate_batch
    call_count = {"calls": 0}

    def flaky_deduplicate_batch(*args, **kwargs):
        call_count["calls"] += 1
        if call_count["calls"] == 1:
            # Simulate a stale connection / idle-in-transaction timeout error on first attempt
            save_db = kwargs.get("db") or (args[0] if args else None)
            if save_db:
                try:
                    save_db.connection().invalidate()
                except Exception:
                    pass
            raise OperationalError(
                "SELECT questions.id FROM questions WHERE content_hash IN (...) [idle-in-transaction timeout]",
                params={},
                orig=Exception("terminating connection due to idle-in-transaction timeout")
            )
        return real_deduplicate_batch(*args, **kwargs)

    pipe.deduplicator.deduplicate_batch = flaky_deduplicate_batch

    # Run the pipeline - should catch OperationalError on attempt 0, retry on attempt 1, and succeed
    asyncio.run(pipe.run(doc_id))

    # Verify that flaky_deduplicate_batch was called twice (initial failed attempt + successful retry)
    assert call_count["calls"] == 2, f"Expected 2 attempts, but got {call_count['calls']}"

    # Verify pipeline status
    reg = INGESTION_STATUS_REGISTRY[doc_id]
    assert reg["status"] == "COMPLETED", f"Pipeline failed: {reg.get('error')}"
    assert reg["questions_extracted"] == 2

    # Verify database state
    verify_db = SessionLocal()
    try:
        saved_doc = verify_db.query(Document).filter(Document.id == doc_id).first()
        assert saved_doc is not None
        assert saved_doc.processing_status == "COMPLETED"
        assert saved_doc.question_count == 2

        saved_questions = verify_db.query(Question).filter(Question.document_id == doc_id).all()
        assert len(saved_questions) == 2
        for q in saved_questions:
            assert tag in q.question_text
    finally:
        verify_db.close()
        if dest.exists():
            dest.unlink()
