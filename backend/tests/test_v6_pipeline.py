"""End to end: upload -> deck details, numbering, v6 topic paths -> Repeated tag across decks."""
import asyncio
import sys
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from backend.app.main import app  # noqa: F401  (registers every model with Base)
from backend.app.config import settings
from backend.app.database import SessionLocal
from backend.app.models.document import Document
from backend.app.models.qshala_tag import QshalaQuestionMetadata, QshalaTag
from backend.app.models.question import Question
from backend.app.models.question_meta import (
    DeckMetadata, QuestionEditHistory, QuestionRepeat, QuestionTopicPath,
)
from backend.app.services.ingestion import pipeline as pipeline_mod
from backend.app.services.ingestion.cancellation import purge_document
from backend.app.services.ingestion.pipeline import IngestionPipeline, INGESTION_STATUS_REGISTRY


class FakeGemini:
    """Stands in for Gemini: two topic paths, a place and a year."""
    calls = 0

    async def generate_json_schema(self, prompt, schema, system_instruction=None, model=None,
                                   thinking_level=None, temperature=0.0):
        FakeGemini.calls += 1
        data = {
            "topics": [
                {"subtopic_l2": "Modern History", "keywords": ["Captain Cook", "Endeavour", "Botany Bay"]},
                {"subtopic_l2": "Indian Politics", "keywords": ["Parliament", "Coalition"]},
            ],
            "geo_location": "Sydney, Australia",
            "india_international": "INTL",
            "relevant_year": 1770,
            "fit": "good",
            "search_summary": "A question about early British exploration.",
        }
        return data, {"in": 10, "out": 10, "think": 0, "cached": 0}


def _build_deck(path: Path, tag: str, n: int = 4):
    prs = Presentation()
    blank = prs.slide_layouts[6]

    def add(text):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
        tb.text_frame.text = text

    add(f"QShala Quest {tag}")
    add("ROUND 1\nConnect questions")
    for i in range(1, n + 1):
        if i == 3:
            add("ROUND 2\nDirect questions")
        add(f"QUESTION {i}\n*Which unique {tag} explorer number {i} charted coast {tag}{i}?")
        add(f"ANSWER {i}\nAnswer: Explorer {tag} {i}\n\nExplanation & Context:\nExplorer {i} sailed in {tag}.")
    prs.save(str(path))


def _ingest(filename: str, tag: str):
    dest = settings.UPLOAD_DIR / f"{uuid.uuid4().hex}.pptx"
    settings.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    _build_deck(dest, tag)
    doc_id = str(uuid.uuid4())
    db = SessionLocal()
    try:
        db.add(Document(id=doc_id, filename=filename, title=filename, year=2022, file_type="pptx",
                        storage_path=str(dest), file_size_bytes=dest.stat().st_size, processing_status="PROCESSING"))
        db.commit()
    finally:
        db.close()
    pipe = IngestionPipeline()
    pipe.qshala_tagger.provider = FakeGemini()
    asyncio.run(pipe.run(doc_id))
    return doc_id


def test_upload_fills_every_v6_tag_and_repeated_tracks_second_deck():
    assert pipeline_mod.USE_QSHALA_TAGGER is True
    tag = uuid.uuid4().hex[:8]
    doc1 = doc2 = None
    try:
        doc1 = _ingest("2022-05-10 PES University Bengaluru Prelims ZVQ History Quiz English.pptx", tag)
        reg = INGESTION_STATUS_REGISTRY[doc1]
        assert reg["status"] == "COMPLETED", reg

        db = SessionLocal()
        deck = db.query(DeckMetadata).filter(DeckMetadata.document_id == doc1).one()
        assert (deck.college, deck.city, deck.round_stage) == ("PES University", "Bengaluru", "Prelims")
        assert deck.quiz_name == "ZVQ History Quiz" and deck.is_varsity is True and deck.deck_year == 2022

        qs = db.query(Question).filter(Question.document_id == doc1).order_by(Question.source_slide_range).all()
        assert len(qs) == 4
        metas = {m.question_id: m for m in db.query(QshalaQuestionMetadata).filter(
            QshalaQuestionMetadata.question_id.in_([q.id for q in qs])).all()}
        assert len(metas) == 4
        nums = sorted((m.question_number, m.round_number, m.question_in_round) for m in metas.values())
        assert nums == [(1, 1, 1), (2, 1, 2), (3, 2, 1), (4, 2, 2)]
        for m in metas.values():
            assert m.star_marked is True                      # the "*" in the slide text
            assert m.is_varsity is True and m.city_name == "Bengaluru"
            assert m.college_institution == "PES University" and m.primary_audience == "College"
            assert m.india_international == "INTL" and m.geo_region == "Sydney, Australia"
            assert m.relevant_year == 1770
            assert m.answer_type == "TEXT" and m.language == "English"
        by_num = {m.question_number: m for m in metas.values()}
        assert by_num[1].question_format == "CONNECT" and by_num[3].question_format == "QA"

        for q in qs:
            assert not q.question_text.startswith("*")        # star stripped so it hashes the same everywhere
            paths = db.query(QuestionTopicPath).filter(QuestionTopicPath.question_id == q.id).order_by(
                QuestionTopicPath.position).all()
            assert [(p.l1, p.l2) for p in paths] == [("History", "Modern History"), ("Politics", "Indian Politics")]
            assert paths[0].keywords[0] == "Captain Cook"
            assert q.topic == "History" and q.subtopic == "Modern History"
            assert q.topics == ["History", "Politics"] and "Captain Cook" in q.tags
            assert db.query(QshalaTag).filter(QshalaTag.question_id == q.id).count() == 1
            assert db.query(QuestionRepeat).filter(QuestionRepeat.question_id == q.id).count() == 0   # first upload: empty
        db.close()

        # Second quiz contains the same questions again.
        doc2 = _ingest("2023-02-11 IIM Bangalore Bengaluru Finals Open Quiz.pptx", tag)
        reg2 = INGESTION_STATUS_REGISTRY[doc2]
        assert reg2["status"] == "COMPLETED", reg2

        db = SessionLocal()
        assert db.query(Question).filter(Question.document_id == doc2).count() == 0   # not stored again
        qs = db.query(Question).filter(Question.document_id == doc1).all()
        for q in qs:
            reps = db.query(QuestionRepeat).filter(QuestionRepeat.question_id == q.id).all()
            assert len(reps) == 1 and reps[0].document_id == doc2
            assert reps[0].quiz_name == "Open Quiz"
            assert reps[0].question_number and reps[0].round_number
            assert q.occurrence_count == 2
            assert db.query(QuestionEditHistory).filter(
                QuestionEditHistory.question_id == q.id, QuestionEditHistory.kind == "REPEAT").count() == 1
        labels = sorted(f"R{r.round_number}Q{r.question_in_round}" for r in db.query(QuestionRepeat).filter(
            QuestionRepeat.document_id == doc2).all())
        assert labels == ["R1Q1", "R1Q2", "R2Q1", "R2Q2"]
        db.close()

        # Running the same deck a third time must not add duplicate Repeated entries for the same deck.
        db = SessionLocal()
        from backend.app.services.question_meta_service import add_repeat
        q0 = db.query(Question).filter(Question.document_id == doc1).first()
        assert add_repeat(db, q0, doc2, {"question_number": 1, "round_number": 1, "question_in_round": 1}, None) is None
        db.close()

        # Deleting the second deck removes its Repeated entries and restores the count.
        db = SessionLocal()
        purge_document(db, doc2)
        doc2 = None
        for q in db.query(Question).filter(Question.document_id == doc1).all():
            db.refresh(q)
            assert db.query(QuestionRepeat).filter(QuestionRepeat.question_id == q.id).count() == 0
            assert q.occurrence_count == 1
        db.close()
    finally:
        db = SessionLocal()
        for d in (doc1, doc2):
            if d:
                purge_document(db, d)
        db.close()


def test_tagger_failure_still_stores_question_as_uncategorised():
    class Broken:
        async def generate_json_schema(self, *a, **k):
            raise ValueError("GEMINI_API_KEY is not configured")

    tag = uuid.uuid4().hex[:8]
    dest = settings.UPLOAD_DIR / f"{uuid.uuid4().hex}.pptx"
    settings.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    _build_deck(dest, tag, n=2)
    doc_id = str(uuid.uuid4())
    db = SessionLocal()
    db.add(Document(id=doc_id, filename="2022 Pune Prelims Quiz.pptx", title="t", year=2022, file_type="pptx",
                    storage_path=str(dest), file_size_bytes=dest.stat().st_size, processing_status="PROCESSING"))
    db.commit()
    db.close()
    try:
        pipe = IngestionPipeline()
        pipe.qshala_tagger.provider = Broken()
        asyncio.run(pipe.run(doc_id))
        assert INGESTION_STATUS_REGISTRY[doc_id]["status"] == "COMPLETED"
        db = SessionLocal()
        qs = db.query(Question).filter(Question.document_id == doc_id).all()
        assert len(qs) == 2 and all(q.topic == "Uncategorised" for q in qs)
        assert db.query(QuestionTopicPath).filter(QuestionTopicPath.question_id.in_([q.id for q in qs])).count() == 0
        db.close()
    finally:
        db = SessionLocal()
        purge_document(db, doc_id)
        db.close()
