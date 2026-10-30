"""
test_qa_tournament_overhaul.py
------------------------------
End-to-End QA Engineer Test Suite for Tournament Ingestion & Question Formats:
1. Verifies all 7 ZVQ tournament presentations are ingested with COMPLETED status.
2. Verifies question format detection:
   - ARRANGE_ORDER (Ticker Tangle): prompt as question_text, clean choices, ordered answer, 60s timer flag.
   - STEPPED_CLUE (Portfolio Builder): 3 progressive clues (+20/-15, +15/-10, +10/-5).
   - CONNECT_4 (Connect 4): multi-clue connect format.
   - BUZZER (Closing Bell / Fin Quiz Buzzer).
3. Verifies taxonomy & curation hygiene:
   - 0 tournament questions with topic "Uncategorised".
   - 0 tournament topic paths with missing L2.
   - 0 tournament questions in draft / Needs Review state.
4. Verifies PowerPoint export:
   - ARRANGE_ORDER / timer slides include top-right timer badge.
"""

import os
import io
import json
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend.app.models.question import Question
from backend.app.models.document import Document
from backend.app.models.question_meta import QuestionTopicPath
from backend.app.models.qshala_tag import QshalaTag, QshalaQuestionMetadata
from backend.app.services.export.exporter import QuizExporter
from pptx import Presentation


@pytest.fixture(scope="module")
def db_session():
    """Connect to the current working SQLite database."""
    db_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "qshala.db"))
    if not os.path.exists(db_path):
        pytest.skip(f"Database file not found at {db_path}")

    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


def test_qa_tournament_documents_completed(db_session):
    """Verify all 7 ZVQ tournament decks exist and are COMPLETED."""
    docs = db_session.query(Document).filter(Document.filename.like("%ZVQ%")).all()
    assert len(docs) >= 7, f"Expected at least 7 ZVQ tournament documents, found {len(docs)}"

    for doc in docs:
        assert doc.processing_status == "COMPLETED", f"Document {doc.filename} has status {doc.processing_status}"
        assert doc.slide_count > 0, f"Document {doc.filename} has invalid slide_count: {doc.slide_count}"

    # Verify Finals decks have 28 competition questions
    finals_docs = [d for d in docs if "Finals" in d.filename or "National" in d.filename]
    for fdoc in finals_docs:
        comp_questions = (
            db_session.query(Question)
            .filter(Question.document_id == fdoc.id, Question.round_number.isnot(None))
            .all()
        )
        assert len(comp_questions) >= 28, (
            f"Finals deck '{fdoc.filename}' expected at least 28 questions, got {len(comp_questions)}"
        )


def test_qa_arrange_order_ticker_tangle(db_session):
    """
    Verify ARRANGE_ORDER questions:
    - Format is ARRANGE_ORDER
    - question_text has prompt instruction (e.g. 'lowest to highest', 'earliest to latest')
    - ordered sequence in answer
    - has_timer is True and timer_seconds == 60 in format_payload
    - round_number is 2
    """
    arrange_qs = db_session.query(Question).filter(Question.question_type == "ARRANGE_ORDER").all()
    assert len(arrange_qs) >= 20, f"Expected at least 20 ARRANGE_ORDER questions across tournament corpus, got {len(arrange_qs)}"

    for q in arrange_qs:
        assert q.question_text and len(q.question_text.strip()) > 5, f"Question {q.id} missing prompt question_text"
        assert q.answer and len(q.answer.strip()) > 0, f"Question {q.id} has empty answer"
        assert q.round_number == 2, f"Ticker Tangle questions must be in Round 2, got Round {q.round_number}"

        payload = q.format_payload or {}
        assert payload.get("has_timer") is True, f"Question {q.id} format_payload should have has_timer == True"
        assert payload.get("timer_seconds") == 60, f"Question {q.id} expected timer_seconds == 60, got {payload.get('timer_seconds')}"


def test_qa_stepped_clue_portfolio_builder(db_session):
    """
    Verify STEPPED_CLUE questions:
    - In Round 3 (Portfolio Builder)
    - format_payload clues list contains 3 progressive clues (+20/-15, +15/-10, +10/-5)
    """
    stepped_qs = db_session.query(Question).filter(Question.question_type == "STEPPED_CLUE").all()
    assert len(stepped_qs) >= 12, f"Expected at least 12 STEPPED_CLUE questions, got {len(stepped_qs)}"

    for q in stepped_qs:
        assert q.round_number == 3, f"Portfolio Builder questions must be in Round 3, got Round {q.round_number}"
        payload = q.format_payload or {}
        clues = payload.get("clues", [])
        assert len(clues) == 3, f"Question {q.id} expected 3 clues in payload, got {len(clues)}"
        pts = [c.get("points") for c in clues]
        assert pts == ["+20/-15", "+15/-10", "+10/-5"], f"Question {q.id} clues have unexpected point values: {pts}"


def test_qa_connect_4_format(db_session):
    """Verify CONNECT_4 questions are properly detected and categorized."""
    connect_qs = db_session.query(Question).filter(Question.question_type == "CONNECT_4").all()
    assert len(connect_qs) >= 15, f"Expected at least 15 CONNECT_4 questions, got {len(connect_qs)}"
    for q in connect_qs:
        assert q.question_type == "CONNECT_4"
        assert q.question_text or (q.question_images and len(q.question_images) > 0), f"Question {q.id} has no text or media"


def test_qa_buzzer_format(db_session):
    """Verify BUZZER questions are properly assigned."""
    buzzer_qs = db_session.query(Question).filter(Question.question_type == "BUZZER").all()
    assert len(buzzer_qs) >= 30, f"Expected at least 30 BUZZER questions, got {len(buzzer_qs)}"
    for q in buzzer_qs:
        assert q.question_type == "BUZZER"


def test_qa_curation_hygiene_zero_uncategorised(db_session):
    """
    Verify curation hygiene:
    - 0 tournament questions with topic == 'Uncategorised'
    - 0 tournament topic paths with missing L2
    - 0 tournament questions in draft / Needs Review state
    """
    tournament_docs = db_session.query(Document).filter(Document.filename.like("%ZVQ%")).all()
    tournament_doc_ids = [d.id for d in tournament_docs]
    assert len(tournament_doc_ids) >= 7

    uncat_qs = (
        db_session.query(Question)
        .filter(
            Question.document_id.in_(tournament_doc_ids),
            Question.topic == "Uncategorised"
        )
        .all()
    )
    assert len(uncat_qs) == 0, f"Found {len(uncat_qs)} tournament questions still Uncategorised: {[q.id for q in uncat_qs[:5]]}"

    uncat_paths = (
        db_session.query(QuestionTopicPath)
        .join(Question, QuestionTopicPath.question_id == Question.id)
        .filter(
            Question.document_id.in_(tournament_doc_ids),
            (QuestionTopicPath.l2.is_(None))
            | (QuestionTopicPath.l2 == "")
            | (QuestionTopicPath.l2 == "Uncategorised")
        )
        .all()
    )
    assert len(uncat_paths) == 0, f"Found {len(uncat_paths)} tournament topic paths with unassigned L2 subtopics!"

    # Tournament questions must not be pending in draft status
    draft_tournament_tags = (
        db_session.query(QshalaTag)
        .join(Question, QshalaTag.question_id == Question.id)
        .filter(
            Question.document_id.in_(tournament_doc_ids),
            QshalaTag.status == "draft"
        )
        .all()
    )
    assert len(draft_tournament_tags) == 0, f"Found {len(draft_tournament_tags)} tournament tags in draft status!"


def test_qa_pptx_export_timer_badge(db_session):
    """
    Verify PowerPoint export:
    - Generates slides for ARRANGE_ORDER questions with the top-right timer badge shape.
    - Timer badge text is '60s TIMER' positioned at top right.
    """
    arrange_q = (
        db_session.query(Question)
        .filter(Question.question_type == "ARRANGE_ORDER")
        .first()
    )
    assert arrange_q is not None, "Need at least one ARRANGE_ORDER question for export test"

    quiz_dict = {
        "title": "QA Tournament Verification Quiz",
        "topic": "Finance & Business",
        "questions": [
            {
                "id": arrange_q.id,
                "question_text": arrange_q.question_text,
                "answer": arrange_q.answer,
                "explanation": arrange_q.explanation,
                "question_format": "ARRANGE_ORDER",
                "format_payload": arrange_q.format_payload,
                "image_refs": arrange_q.image_refs or [],
                "question_images": arrange_q.question_images or [],
                "answer_images": arrange_q.answer_images or [],
                "is_approved": True,
            }
        ],
    }

    pptx_bytes = QuizExporter.to_pptx(quiz_dict, quiz_type="corporate")
    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 1000, "PPTX export byte size too small"

    prs = Presentation(io.BytesIO(pptx_bytes))
    found_timer_shape = False

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and "60s TIMER" in shape.text_frame.text:
                found_timer_shape = True
                left_in = shape.left.inches
                top_in = shape.top.inches
                assert left_in > 9.0, f"Timer badge not positioned on right side: left={left_in}"
                assert top_in < 1.0, f"Timer badge not positioned at top: top={top_in}"

    assert found_timer_shape, "PPTX export did not produce top-right '60s TIMER' badge for ARRANGE_ORDER question"
