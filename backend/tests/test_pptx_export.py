import io
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.dml.color import RGBColor
from backend.app.main import app
from backend.app.database import SessionLocal
from backend.app.models.quiz import Quiz, GeneratedQuestion
from backend.app.services.export.exporter import QuizExporter, _clean_quiz_title

client = TestClient(app)


def test_clean_quiz_title_helper():
    """Verify raw prompt text is converted to crisp title strings."""
    assert _clean_quiz_title("give me 7 questions on government policy and make a quiz", "Finance") == "Government Policy Quiz"
    assert _clean_quiz_title("generate 10 questions about Indian Space Program", "Science") == "Indian Space Program Quiz"
    assert _clean_quiz_title("Bangalore South Zone", "Finance") == "Bangalore South Zone"
    assert _clean_quiz_title("", "Astronomy") == "Astronomy Quiz"


def test_pptx_export_bengaluru_south_theme():
    """Verify export produces a valid 16:9 presentation with authentic Bengaluru South tournament slides."""
    sample_quiz = {
        "id": "test-bengaluru-quiz",
        "title": "Bengaluru South Open Championship",
        "topic": "Karnataka Business & Culture",
        "difficulty": "Medium",
        "questions": [
            {
                "id": "q1",
                "question_text": "Which plant was adopted by Karnataka Bank in 2002 to brand its ATM network?",
                "options": ["A) Money plant", "B) Bamboo", "C) Tulsi", "D) Banyan"],
                "answer": "Money plant",
                "explanation": "Epipremnum aureum, commonly known as money plant, symbolizes prosperity.",
                "difficulty": "Medium",
                "is_approved": True,
            },
            {
                "id": "q2",
                "question_text": "Which iconic IT company was founded in Pune in 1981 but later moved its headquarters to Electronics City, Bengaluru?",
                "options": ["A) Infosys", "B) Wipro", "C) TCS", "D) Mindtree"],
                "answer": "Infosys",
                "explanation": "Infosys relocated to Bengaluru in 1983 and established its landmark campus in Electronics City.",
                "difficulty": "Easy",
                "is_approved": True,
            }
        ]
    }

    # Test direct QuizExporter.to_pptx with Bengaluru South theme and social slides
    pptx_bytes = QuizExporter.to_pptx(sample_quiz, theme="bengaluru_south", include_social_slides=True)
    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 1000

    prs = Presentation(io.BytesIO(pptx_bytes))
    # 4 intro slides (Title + Social + QR + Round) + 2 * 2 (Q+A) = 8 slides
    assert len(prs.slides) == 8
    assert prs.slide_width.inches == pytest.approx(13.333, 0.01)
    assert prs.slide_height.inches == pytest.approx(7.5, 0.01)

    # 1. Slide 1 (Title): Verify dark font color (NOT white) and custom/cleaned text
    title_slide = prs.slides[0]
    p_title = title_slide.shapes[1].text_frame.paragraphs[0]
    assert "Bengaluru South Open Championship" in p_title.text
    # Dark charcoal #0F172A: RGB(15, 23, 42)
    assert p_title.font.color.rgb == RGBColor(15, 23, 42)

    # 2. Slide 2: Social Media
    s2 = prs.slides[1]
    s2_texts = [p.text for s in s2.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Social Media" in t for t in s2_texts)
    assert any("#varsityquiz" in t for t in s2_texts)

    # 3. Slide 3: QR Codes (Follow @Zerodhavarsity)
    s3 = prs.slides[2]
    s3_texts = [p.text for s in s3.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Get Curious & Stay Curious" in t for t in s3_texts)
    # Check that pictures (QR codes) were embedded
    pics = [s for s in s3.shapes if s.shape_type == 13]  # 13 = PICTURE
    assert len(pics) >= 3

    # 4. Slide 4: Round Divider
    s4 = prs.slides[3]
    s4_texts = [p.text for s in s4.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Preliminary Round" in t for t in s4_texts)

    # 5. Slide 5: Question 1 (badge "1." dark, text dark)
    q1_slide = prs.slides[4]
    q1_texts = [p.text for s in q1_slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Karnataka Bank" in t for t in q1_texts)
    assert any("1." in t for t in q1_texts)

    # 6. Slide 6: Answer 1 (emerald green reveal)
    ans1_slide = prs.slides[5]
    ans1_texts = [p.text for s in ans1_slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Money plant" in t for t in ans1_texts)


def test_pptx_export_custom_title_and_no_social_toggle():
    """Verify custom title/subtitle parameters and social media toggle."""
    sample_quiz = {
        "id": "test-custom-title",
        "title": "give me 7 questions on government policy and make a quiz",
        "topic": "Finance",
        "questions": [
            {
                "question_text": "What does SEBI stand for?",
                "answer": "Securities and Exchange Board of India",
                "is_approved": True,
            }
        ]
    }

    # Without custom title, it cleans the prompt
    pptx_bytes_auto = QuizExporter.to_pptx(sample_quiz, theme="bengaluru_south", include_social_slides=False)
    prs_auto = Presentation(io.BytesIO(pptx_bytes_auto))
    # 1 title slide + 1 question + 1 answer = 3 slides (no social slides)
    assert len(prs_auto.slides) == 3
    t1_texts = [p.text for s in prs_auto.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("Government Policy Quiz" in t for t in t1_texts)

    # With explicit custom title and subtitle
    pptx_bytes_custom = QuizExporter.to_pptx(
        sample_quiz,
        theme="bengaluru_south",
        custom_title="National Finance Championship 2026",
        custom_subtitle="Bangalore Campus Edition - Prelims",
        include_social_slides=False
    )
    prs_custom = Presentation(io.BytesIO(pptx_bytes_custom))
    t2_texts = [p.text for s in prs_custom.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert any("National Finance Championship 2026" in t for t in t2_texts)
    assert any("Bangalore Campus Edition - Prelims" in t for t in t2_texts)


def test_pptx_export_minimalist_fallback():
    """Verify fallback to minimalist styling when requested."""
    sample_quiz = {
        "id": "test-minimalist",
        "title": "Minimalist Quiz",
        "topic": "General",
        "questions": [
            {
                "id": "q1",
                "question_text": "What is the capital of India?",
                "answer": "New Delhi",
                "explanation": "New Delhi was inaugurated in 1931.",
                "is_approved": True,
            }
        ]
    }
    pptx_bytes = QuizExporter.to_pptx(sample_quiz, theme="minimalist")
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 3


def test_pptx_export_api_endpoints():
    """Test both GET /quizzes/{id}/export?format=pptx and POST /quizzes/export?format=pptx with custom parameters."""
    db = SessionLocal()
    quiz = None
    try:
        # Create a test quiz in database
        quiz = Quiz(
            title="API Test Quiz",
            topic="History",
            difficulty="Balanced",
            quiz_format="TOURNAMENT_PRELIMS"
        )
        db.add(quiz)
        db.commit()
        db.refresh(quiz)

        gq = GeneratedQuestion(
            quiz_id=quiz.id,
            question_text="Who was the architect of New Delhi?",
            answer="Sir Edwin Lutyens",
            explanation="Lutyens planned New Delhi with Sir Herbert Baker.",
            order_index=1,
            is_approved=True
        )
        db.add(gq)
        db.commit()

        # Test GET with custom parameters
        url = (
            f"/api/v1/quizzes/{quiz.id}/export?format=pptx&theme=bengaluru_south"
            f"&custom_title=Campus+Arena+Quiz&custom_subtitle=Zonal+Prelims&include_social_slides=true"
        )
        res_get = client.get(url)
        assert res_get.status_code == 200
        assert "presentationml" in res_get.headers["content-type"]
        assert len(res_get.content) > 1000
        prs_get = Presentation(io.BytesIO(res_get.content))
        # 4 intro slides + 2 (Q+A) = 6 slides
        assert len(prs_get.slides) == 6
        get_title_texts = [p.text for s in prs_get.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs]
        assert any("Campus Arena Quiz" in t for t in get_title_texts)
        assert any("Zonal Prelims" in t for t in get_title_texts)

        # Test POST /api/v1/quizzes/export?format=pptx
        payload = {
            "title": "give me 5 questions on physics",
            "topic": "Science",
            "questions": [
                {
                    "question_text": "What is the speed of light in vacuum?",
                    "answer": "Approximately 300,000 km/s",
                    "explanation": "299,792,458 metres per second precisely.",
                    "is_approved": True
                }
            ]
        }
        res_post = client.post(
            "/api/v1/quizzes/export?format=pptx&theme=bengaluru_south&include_social_slides=false",
            json=payload
        )
        assert res_post.status_code == 200
        assert len(res_post.content) > 1000
        prs_post = Presentation(io.BytesIO(res_post.content))
        # 1 title slide + 1 question + 1 answer = 3 slides
        assert len(prs_post.slides) == 3
        post_title_texts = [p.text for s in prs_post.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs]
        assert any("Physics Quiz" in t for t in post_title_texts)

    finally:
        if quiz:
            db.query(GeneratedQuestion).filter(GeneratedQuestion.quiz_id == quiz.id).delete()
            db.query(Quiz).filter(Quiz.id == quiz.id).delete()
            db.commit()
        db.close()
