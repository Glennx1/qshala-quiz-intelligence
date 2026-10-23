"""The parser no longer holds the whole deck's media in memory, and uploads media in parallel."""
import io
import sys
import uuid
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from backend.app.config import settings
from backend.app.services.ingestion.pptx_parser import PPTXParser


class FakeR2:
    public_base_url = "https://cdn.example.com"

    def __init__(self):
        self.keys = []

    def upload_bytes(self, key, data, content_type=None):
        self.keys.append(key)
        return f"{self.public_base_url}/{key}"


def _deck(path: Path):
    prs = Presentation()
    for n in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        img = Image.effect_noise((600, 400), 60 + n * 30).convert("RGB")
        buf = io.BytesIO()
        img.save(buf, "PNG")
        png = path.parent / f"i{n}.png"
        png.write_bytes(buf.getvalue())
        s.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(4))
        s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = f"QUESTION {n + 1}"
    prs.save(str(path))


def test_images_found_without_holding_bytes_and_r2_uploads_finish(tmp_path, monkeypatch):
    fake = FakeR2()
    import backend.app.services.storage.r2_storage as r2mod
    monkeypatch.setattr(r2mod, "get_media_storage", lambda: fake)
    deck = tmp_path / "d.pptx"
    _deck(deck)
    doc = f"memtest-{uuid.uuid4().hex[:8]}"
    slides = PPTXParser(str(deck), doc).parse()
    try:
        urls = [u for s in slides for u in s["image_paths"]]
        assert len(urls) == 3 and all(u.startswith("https://cdn.example.com/slides/") for u in urls)
        assert len(fake.keys) >= 3                                  # every upload had finished when parse returned
        for s in slides:
            for item in s["all_media_items"]:
                assert "raw_bytes" not in item                      # nothing keeps the deck's bytes alive
    finally:
        import shutil
        shutil.rmtree(settings.SLIDES_DIR / doc, ignore_errors=True)
