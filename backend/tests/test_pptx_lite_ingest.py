"""PPTX ingestion on a small memory budget.

The parser reads slide structure from a "lite" copy of the deck (media replaced by placeholders),
streams each distinct media file once from the original zip to disk, uploads to R2 from that file,
and keeps template art (layout/master/theme backgrounds) off slide 1.
"""
import hashlib
import io
import re
import shutil
import sys
import uuid
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from backend.app.config import settings
from backend.app.services.ingestion.pptx_parser import PPTXParser, build_lite_package


def _png(seed: int, size=(640, 420)) -> bytes:
    img = Image.effect_noise(size, 40 + seed * 25).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _md5(b: bytes) -> str:
    return hashlib.md5(b).hexdigest()


@pytest.fixture
def deck(tmp_path):
    """4 slides. Slides 1-2 show the SAME picture, slide 3 a different one, slide 4 only text.
    Also carries a layout background picture, an orphan picture and an audio clip."""
    pics = {n: _png(n) for n in range(2)}
    prs = Presentation()
    for idx, pic_no in enumerate([0, 0, 1, None]):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = f"QUESTION {idx + 1}"
        if pic_no is not None:
            p = tmp_path / f"p{pic_no}.png"
            p.write_bytes(pics[pic_no])
            s.shapes.add_picture(str(p), Inches(1), Inches(1), Inches(4))
    base = tmp_path / "base.pptx"
    prs.save(str(base))

    layout_art, orphan, audio = _png(7, (900, 500)), _png(8, (300, 300)), b"ID3" + bytes(range(256)) * 40
    out = tmp_path / "deck.pptx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "[Content_Types].xml":
                text = data.decode()
                if 'Extension="mp3"' not in text:
                    text = text.replace("<Default ", '<Default Extension="mp3" ContentType="audio/mpeg"/><Default ', 1)
                data = text.encode()
            elif info.filename == "ppt/slideLayouts/_rels/slideLayout7.xml.rels":
                text = data.decode().replace(
                    "</Relationships>",
                    '<Relationship Id="rIdArt" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                    'Target="../media/layout_art.png"/></Relationships>')
                data = text.encode()
            elif info.filename == "ppt/slides/_rels/slide2.xml.rels":
                text = data.decode().replace(
                    "</Relationships>",
                    '<Relationship Id="rIdAud" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio" '
                    'Target="../media/clip.mp3"/></Relationships>')
                data = text.encode()
            zout.writestr(info, data)
        zout.writestr("ppt/media/layout_art.png", layout_art)
        zout.writestr("ppt/media/orphan.png", orphan)
        zout.writestr("ppt/media/clip.mp3", audio)
    return {"path": out, "pics": pics, "layout_art": layout_art, "orphan": orphan, "audio": audio}


@pytest.fixture
def doc_id():
    doc = f"lite-{uuid.uuid4().hex[:8]}"
    yield doc
    shutil.rmtree(settings.SLIDES_DIR / doc, ignore_errors=True)
    # the parser also mirrors media into the Next.js public folder when it exists (dev setups)
    shutil.rmtree(Path(__file__).resolve().parents[2] / "frontend" / "public" / "static" / "slides" / doc, ignore_errors=True)


@pytest.fixture(autouse=True)
def _no_r2_by_default(monkeypatch):
    import backend.app.services.storage.r2_storage as r2mod
    monkeypatch.setattr(r2mod, "get_media_storage", lambda: None)


def _media_dir(doc):
    return settings.SLIDES_DIR / doc / "media"


def test_lite_package_drops_media_but_keeps_structure(deck):
    lite = build_lite_package(str(deck["path"]))
    assert lite.getbuffer().nbytes < deck["path"].stat().st_size       # media is gone
    with zipfile.ZipFile(lite) as z:
        assert z.read("ppt/media/layout_art.png") == b"\x00"
        assert b"QUESTION 1" in z.read("ppt/slides/slide1.xml")
    assert len(Presentation(build_lite_package(str(deck["path"]))).slides) == 4


def test_shared_picture_is_stored_once_and_bytes_are_intact(deck, doc_id):
    slides = PPTXParser(str(deck["path"]), doc_id).parse()

    # (the orphan picture nothing points at is filed under slide 1, as before)
    slide1 = [u for u in slides[0]["image_paths"] if not u.endswith("orphan.png")]
    assert slide1 == slides[1]["image_paths"]                           # same picture -> same file
    assert len(slide1) == 1 and len(slides[2]["image_paths"]) == 1
    assert slides[3]["image_paths"] == []
    assert slide1 != slides[2]["image_paths"]

    files = {p.name: _md5(p.read_bytes()) for p in _media_dir(doc_id).iterdir()}
    for slide_no, pic_no in [(0, 0), (2, 1)]:
        name = slides[slide_no]["image_paths"][0].rsplit("/", 1)[1]
        assert files[name] == _md5(deck["pics"][pic_no])                # the real bytes, not the placeholder
    assert not any(p.suffix == ".part" for p in _media_dir(doc_id).iterdir())


def test_audio_is_saved_with_original_bytes(deck, doc_id):
    slides = PPTXParser(str(deck["path"]), doc_id).parse()
    assert slides[1]["has_audio"] and len(slides[1]["audio_paths"]) == 1
    saved = _media_dir(doc_id) / slides[1]["audio_paths"][0].rsplit("/", 1)[1]
    assert saved.read_bytes() == deck["audio"]


def test_template_art_stays_off_slide_one_by_default(deck, doc_id, monkeypatch):
    monkeypatch.delenv("PPTX_ATTACH_TEMPLATE_MEDIA", raising=False)
    slides = PPTXParser(str(deck["path"]), doc_id).parse()
    names = {Path(m["rel_url"]).name for m in slides[0]["all_media_items"]}
    assert "layout_art.png" not in names                                # layout background: skipped
    assert "orphan.png" in names                                        # media nothing points at: still kept
    assert not (_media_dir(doc_id) / "layout_art.png").exists()


def test_template_art_flag_restores_old_behaviour(deck, doc_id, monkeypatch):
    monkeypatch.setenv("PPTX_ATTACH_TEMPLATE_MEDIA", "1")
    slides = PPTXParser(str(deck["path"]), doc_id).parse()
    names = {Path(m["rel_url"]).name for m in slides[0]["all_media_items"]}
    assert {"layout_art.png", "orphan.png"} <= names


class FileFakeR2:
    """Records what it is asked to upload, reading the file at upload time like boto3 would."""
    public_base_url = "https://cdn.example.com"

    def __init__(self):
        self.uploads = {}

    def upload_file(self, key, path, content_type=None):
        self.uploads[key] = _md5(Path(path).read_bytes())
        return f"{self.public_base_url}/{key}"


def test_r2_upload_reads_from_disk_and_urls_point_at_r2(deck, doc_id, monkeypatch):
    fake = FileFakeR2()
    import backend.app.services.storage.r2_storage as r2mod
    monkeypatch.setattr(r2mod, "get_media_storage", lambda: fake)
    slides = PPTXParser(str(deck["path"]), doc_id).parse()

    url0 = slides[0]["image_paths"][0]
    assert url0.startswith("https://cdn.example.com/slides/")
    key0 = url0[len("https://cdn.example.com/"):]
    assert fake.uploads[key0] == _md5(deck["pics"][0])
    assert len(fake.uploads) == len(set(fake.uploads))                  # shared picture uploaded once


def test_without_local_copy_files_are_staged_uploaded_and_cleaned_up(deck, doc_id, monkeypatch, tmp_path):
    fake = FileFakeR2()
    import backend.app.services.storage.r2_storage as r2mod
    monkeypatch.setattr(r2mod, "get_media_storage", lambda: fake)
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path / "stage"), raising=False)
    (tmp_path / "stage").mkdir()

    parser = PPTXParser(str(deck["path"]), doc_id, write_to_disk=False)
    slides = parser.parse()

    assert slides[0]["image_paths"][0].startswith("https://cdn.example.com/")
    assert not (settings.SLIDES_DIR / doc_id / "media").exists()        # nothing written locally
    assert all(m["local_path"] is None for s in slides for m in s["all_media_items"])
    assert any(_md5(deck["pics"][0]) == h for h in fake.uploads.values())
    assert list((tmp_path / "stage").iterdir()) == []                   # staging area removed


def test_no_local_copy_and_no_r2_writes_nothing(deck, doc_id):
    slides = PPTXParser(str(deck["path"]), doc_id, write_to_disk=False).parse()
    assert slides[0]["has_images"]
    assert not (settings.SLIDES_DIR / doc_id / "media").exists()


def test_failed_upload_does_not_fail_the_parse(deck, doc_id, monkeypatch):
    class Boom(FileFakeR2):
        def upload_file(self, key, path, content_type=None):
            raise RuntimeError("R2 down")
    import backend.app.services.storage.r2_storage as r2mod
    monkeypatch.setattr(r2mod, "get_media_storage", lambda: Boom())
    monkeypatch.setattr(PPTXParser, "_upload_with_retry", staticmethod(
        lambda r2, key, path, attempts=1: r2.upload_file(key, path)))
    slides = PPTXParser(str(deck["path"]), doc_id).parse()
    assert slides[0]["has_images"]


def test_cancel_stops_parsing(deck, doc_id):
    from backend.app.services.ingestion.cancellation import IngestionCancelled
    parser = PPTXParser(str(deck["path"]), doc_id)
    parser.should_cancel = lambda: True
    with pytest.raises(IngestionCancelled):
        parser.parse()


def test_small_gif_and_tiny_images_are_ignored(tmp_path, doc_id):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tiny = tmp_path / "tiny.png"
    Image.new("RGB", (20, 20), "red").save(tiny)
    s.shapes.add_picture(str(tiny), Inches(1), Inches(1), Inches(1))
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    slides = PPTXParser(str(path), doc_id).parse()
    assert slides[0]["image_paths"] == []


def test_slide_media_follows_the_slide_not_its_position(tmp_path, doc_id):
    """Slides reordered in PowerPoint keep their old file numbers; the zip fallback must follow the slide."""
    prs = Presentation()
    pics = [_png(n + 20) for n in range(2)]
    for n in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = f"Q{n + 1}"
        p = tmp_path / f"s{n}.png"
        p.write_bytes(pics[n])
        s.shapes.add_picture(str(p), Inches(1), Inches(1), Inches(4))
    base = tmp_path / "base.pptx"
    prs.save(str(base))
    # swap the order in presentation.xml: slide2.xml is now first
    out = tmp_path / "swapped.pptx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "ppt/presentation.xml":
                ids = re.findall(rb"<p:sldId [^>]*/>", data)
                assert len(ids) == 2
                data = data.replace(ids[0] + ids[1], ids[1] + ids[0])
            zout.writestr(info, data)
    slides = PPTXParser(str(out), doc_id).parse()
    assert slides[0]["extracted_text"] == "Q2"
    first_img = _media_dir(doc_id) / slides[0]["image_paths"][0].rsplit("/", 1)[1]
    assert _md5(first_img.read_bytes()) == _md5(pics[1])                # Q2's own picture
    assert len(slides[0]["image_paths"]) == 1


def test_r2_service_upload_file_streams_from_disk(tmp_path, monkeypatch):
    boto3 = pytest.importorskip("boto3")
    moto = pytest.importorskip("moto")
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "x")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "x")
    from backend.app.services.storage.r2_storage import CloudflareR2Service

    with moto.mock_aws():
        svc = CloudflareR2Service.__new__(CloudflareR2Service)
        svc.bucket_name = "qshala-test"
        svc.public_base_url = "https://cdn.example.com"
        svc.is_configured = True
        svc._client = boto3.client("s3", region_name="us-east-1")
        svc._client.create_bucket(Bucket="qshala-test")

        small = tmp_path / "a.png"
        small.write_bytes(b"\x89PNG" + b"1" * 5000)
        big = tmp_path / "v.mp4"
        big.write_bytes(b"x" * (20 * 1024 * 1024))                      # above the multipart threshold

        assert svc.upload_file("slides/d/media/a.png", str(small)) == "https://cdn.example.com/slides/d/media/a.png"
        svc.upload_file("slides/d/media/v.mp4", str(big))
        got = svc._client.get_object(Bucket="qshala-test", Key="slides/d/media/a.png")
        assert got["ContentType"] == "image/png" and got["Body"].read() == small.read_bytes()
        head = svc._client.head_object(Bucket="qshala-test", Key="slides/d/media/v.mp4")
        assert head["ContentLength"] == big.stat().st_size and head["ContentType"] == "video/mp4"
